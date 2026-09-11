from pathlib import Path
from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(r"C:\Users\Asus\Documents\estabilizacion\ah sports\App de gestion\docs\Guia de carga de datos AH Sports.docx")
OUT = Path(r"C:\Users\Asus\Documents\estabilizacion\ah sports\App de gestion\docs\GUIA-MAESTRA-CARGA-DE-DATOS-AH-SPORTS.docx")
NAVY, PALE, GRID, BLACK, DARK = "1F3A5F", "F4F6F8", "D9D9D9", "000000", "333333"

def set_font(run, size=9.2, bold=False, color=DARK, name="Aptos"):
    run.font.name=name; rpr=run._element.get_or_add_rPr(); rpr.rFonts.set(qn("w:ascii"),name); rpr.rFonts.set(qn("w:hAnsi"),name)
    run.font.size=Pt(size); run.bold=bold; run.font.color.rgb=RGBColor.from_string(color)

def sp(p, after=5, before=0, line=1.1):
    f=p.paragraph_format; f.space_after=Pt(after); f.space_before=Pt(before); f.line_spacing=line

def cell_setup(c, fill=None):
    if fill:
        pr=c._tc.get_or_add_tcPr(); sh=OxmlElement("w:shd"); sh.set(qn("w:fill"),fill); pr.append(sh)
    pr=c._tc.get_or_add_tcPr(); m=OxmlElement("w:tcMar")
    for side in ("top","start","bottom","end"):
        x=OxmlElement("w:"+side); x.set(qn("w:w"),"95"); x.set(qn("w:type"),"dxa"); m.append(x)
    pr.append(m); b=OxmlElement("w:tcBorders")
    for side in ("top","left","bottom","right","insideH","insideV"):
        x=OxmlElement("w:"+side); x.set(qn("w:val"),"single"); x.set(qn("w:sz"),"6"); x.set(qn("w:color"),GRID); b.append(x)
    pr.append(b); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, heads, rows, widths, size=7.8, fill=NAVY):
    t=doc.add_table(rows=1,cols=len(heads)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    tr=t.rows[0]._tr.get_or_add_trPr(); tr.append(OxmlElement("w:tblHeader")); tr.append(OxmlElement("w:cantSplit"))
    for i,h in enumerate(heads):
        c=t.rows[0].cells[i]; c.text=""; cell_setup(c,fill); p=c.paragraphs[0]; sp(p,0,0,1); set_font(p.add_run(str(h)),size,True,"FFFFFF")
    for ri,data in enumerate(rows):
        row=t.add_row(); row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for i,v in enumerate(data):
            c=row.cells[i]; c.text=""; cell_setup(c,PALE if ri%2 else None); p=c.paragraphs[0]; sp(p,0,0,1); set_font(p.add_run(str(v)),size)
    for row in t.rows:
        for i,w in enumerate(widths):
            c=row.cells[i]; c.width=Inches(w); pr=c._tc.get_or_add_tcPr(); tw=OxmlElement("w:tcW"); tw.set(qn("w:w"),str(int(w*1440))); tw.set(qn("w:type"),"dxa"); pr.append(tw)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def heading(doc,text,level=1):
    p=doc.add_paragraph(style=f"Heading {level}"); p.paragraph_format.keep_with_next=True; sp(p,6 if level==1 else 3,14 if level==1 else 8,1); set_font(p.add_run(text),{1:17,2:13,3:10.5}[level],True,BLACK)

def body(doc,text):
    p=doc.add_paragraph(); sp(p,6,0,1.12); set_font(p.add_run(text),9.35)

def bullets(doc,items,checks=False):
    for item in items:
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.45); sp(p,3,0,1.07); set_font(p.add_run(("[ ] " if checks else "• ")+item),9.1)

def note(doc,lead,text):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.4); sp(p,7,3,1.08); set_font(p.add_run(lead+" "),9.25,True); set_font(p.add_run(text),9.25)

def code(doc,text):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.55); sp(p,7,2,1); lines=text.split("\n")
    for i,line in enumerate(lines):
        r=p.add_run(line); set_font(r,8.5,False,DARK,"Consolas")
        if i<len(lines)-1: r.add_break()

def fields(doc,rows): table(doc,["Campo","Qué registrar","Obligatorio","Ejemplo / criterio"],rows,[1.3,2.8,.85,2.3],7.8)

doc=Document(BASE)
# Preserve the practical form already supplied, but make the title explicit and add a master reference section.
if doc.paragraphs:
    p=doc.paragraphs[0]; p.text=""; p.style="Title"; sp(p,12,0,1); set_font(p.add_run("Guía maestra de carga de datos AH Sports"),24,True,BLACK,"Aptos Display")
for sec in doc.sections:
    sec.top_margin=Cm(1.65); sec.bottom_margin=Cm(1.55); sec.left_margin=Cm(1.8); sec.right_margin=Cm(1.8)
    h=sec.header.paragraphs[0]; h.text=""; set_font(h.add_run("AH SPORTS  |  MANUAL OPERATIVO DE DATOS"),8,True,BLACK); sp(h,0,0,1)
    f=sec.footer.paragraphs[0]; f.text=""; set_font(f.add_run("Uso interno operativo  |  Neon + panel administrativo"),8,False,"666666"); sp(f,0,0,1)

doc.add_page_break(); heading(doc,"Parte II. Referencia maestra del sistema")
body(doc,"Esta segunda parte amplía la ficha rellenable inicial y deja documentadas las relaciones exactas del proyecto. Sirve como manual de conversación: recibimos audio o texto, separamos entidades, validamos invariantes y cargamos en orden.")
table(doc,["Sección","Uso"],[["1. Propósito y alcance","Qué se carga y qué no se debe enviar."],["2. Modelo y relaciones","Cómo se conectan proveedores, materiales, productos, recetas y pedidos."],["3. Orden de carga","Dependencias y secuencia operativa."],["4. Proveedores","Ficha de contacto y vínculo con insumos."],["5. Insumos","Unidades, categorías, precios, rendimiento y merma."],["6. Catálogo","Categoría, subcategoría, tipo y familia."],["7. Familias y moldes","Reutilización de moldes sin compartir BOM."],["8. Productos","SKU, garment, bundle, zonas y mínimo."],["9. Talles","Medidas terminadas y consumo por talle."],["10. Recetas BOM","direct/yield, fórmulas y ejemplos."],["11. Técnicas y pricing","Costos, m², setup, margen y redondeo."],["12. Kits","Componentes y mapping same_label/fixed."],["13. Pedidos","Cotización pública, estados, pagos y snapshots."],["14. Política flexible","Pendientes versus invariantes."],["15. Panel y APIs","Dónde cargar y rutas útiles."],["16. Audio","Plantillas para dictar información."],["17. Gobierno","Actualizaciones y trazabilidad."],["18. Estado técnico","Migraciones y verificaciones actuales."],["Apéndice","Registros completos compactos."]],[2.25,4.9],8.0,"5C7897")

heading(doc,"1. Propósito y alcance")
body(doc,"AH Sports convierte conocimiento del taller en datos reutilizables para cotizar, producir y hacer seguimiento. Esta guía explica qué preguntar y cómo relacionar los datos sin exigir conocer tablas o códigos internos.")
body(doc,"El alcance cubre proveedores, materiales, técnicas, catálogo, familias, moldes, productos, talles, recetas/BOM, conjuntos, pedidos, pagos, caja y snapshots. La carga es progresiva: lo que no se sabe se marca pendiente, salvo lo que rompería una relación o un cálculo.")
heading(doc,"1.1 Principios",2); bullets(doc,["Cada proveedor se carga una vez y se selecciona desde los materiales.","Cada material se registra por variante útil y conserva su precio de compra actual.","Cada producto tiene SKU, nombre y clasificación; familia y molde son conceptos separados.","Cada receta enlaza materiales y expresa consumo en unidad de consumo.","Cada bundle agrupa productos garment y no se anida.","Los cambios futuros no reescriben snapshots históricos.","No enviar secretos, tokens, SESSION_SECRET ni DATABASE_URL."])

heading(doc,"2. Modelo general y relaciones")
table(doc,["Entidad","Representa","Se relaciona con","Clave"],[["Proveedor","Vendedor de materiales.","Materiales","Nombre normalizado y contacto."],["Material","Tela, hilo, avío, accesorio o packaging.","Proveedor y BOM","Unidad y precio de compra."],["Técnica","Sublimación, DTF, bordado, etc.","Recetas, líneas, aplicaciones","Unidad, m² y/o setup."],["Catálogo","Clasificación general editable.","Producto","SELECT; no textarea."],["Familia/molde","Estructura y medidas.","Producto y talles","Molde reutilizable; BOM separado."],["Producto","Artículo vendible.","Talles, recetas, pedidos","SKU y productKind."],["Talle","Variante del producto.","BOM y prendas","Medidas terminadas."],["Receta BOM","Consumo y costo.","Producto/material/técnica","direct/yield, merma."],["Bundle","Producto compuesto.","Componentes/talles","same_label/fixed."],["Pedido","Compromiso comercial.","Líneas/pagos/arte","Snapshot inmutable."]],[1.15,2.05,2.1,2.1],7.55)
code(doc,"PROVEEDOR\n  └── 0..n MATERIALES\n                    └── 0..n ÍTEMS BOM\nPRODUCTO ── 0..n RECETAS ── 0..n ÍTEMS ── 1 MATERIAL\n   ├── 0..n TALLES\n   ├── 0..1 MOLDE ── FAMILIA + MEDIDAS\n   └── bundle ── COMPONENTES garment\nPEDIDO ── LÍNEAS ── PRODUCTO ── TALLE/TÉCNICA\n   ├── PAGOS ── CAJA\n   └── SNAPSHOT (no se pisa)")

heading(doc,"3. Orden recomendado de carga")
table(doc,["Paso","Cargar","Por qué"],[["1","Proveedores","Los materiales necesitan una entidad de contacto."],["2","Categorías del catálogo","Los productos deben elegir opciones por SELECT."],["3","Materiales","Las recetas seleccionan materiales existentes."],["4","Técnicas","Las recetas y líneas pueden enlazarlas."],["5","Familias y moldes","Preparan productos textiles y esquemas de medidas."],["6","Productos","Reúnen identidad, clasificación y relaciones."],["7","Talles y medidas","Permiten BOM y pedidos por talle."],["8","Recetas BOM","Conectan producto, material, técnica y costo."],["9","Conjuntos","Agrupan productos ya cotizables."],["10","Pedidos","Usan catálogo maduro y generan snapshots."]],[.45,1.8,5.2],7.8,"5C7897")
note(doc,"Excepción.","Un registro incompleto puede quedar como pendiente. No hay que inventar un dato para completar una pantalla; se pregunta cuando sea necesario para guardar la relación o cotizar.")

heading(doc,"4. Proveedores")
body(doc,"El proveedor es una entidad reutilizable. Textil SA, Textil S.A. y textil sa deben representar el mismo registro, no tres. Si el mismo material se compra a dos proveedores con distinta calidad/precio, registrar variantes claras.")
fields(doc,[["name","Nombre comercial o razón social.","Sí","Textil SA"],["phone","Teléfono o WhatsApp.","No, recomendado","+54 370 4XX XXXX"],["email","Correo comercial.","No, recomendado","ventas@textil.example"],["address","Domicilio o localidad.","No","Av. Central 123, Formosa"],["contactName","Persona habitual.","No","María López"],["taxId","CUIT u otro identificador.","No","30-00000000-0"],["notes","Horarios, mínimos, plazos, condiciones.","No","Despacha martes y jueves"],["active","Disponible para compras nuevas.","No","Desactivar, no borrar histórico"]])
body(doc,"El material guarda supplierId como vínculo canónico. El texto supplier queda por compatibilidad histórica. Cambiar de proveedor afecta compras y cotizaciones futuras, no snapshots.")

heading(doc,"5. Insumos y materiales")
body(doc,"Un material es cualquier insumo comprable que entra al costo: tela, hilo, cierre, botón, elástico, avío, accesorio, etiqueta, bolsa o packaging. Se carga por variante útil y se reutiliza en recetas.")
table(doc,["Unidad","Uso","Cálculo"],[["kilo","Tela/hilo comprado por kg.","unitPrice / metersPerKilo para metro."],["metro","Compra y consumo por metro.","unitPrice ya es por metro."],["unidad","Cierre, botón, etiqueta, bolsa.","cantidad × unitPrice."],["rollo / pack","Solo si la presentación se controla así.","Documentar conversión y contenido."]],[1.0,3.0,3.4],8.1)
note(doc,"Kilo versus metro.","Si unit=kilo, metersPerKilo > 0 es obligatorio. Una tela a $30.000/kg con rendimiento de 5 m/kg cuesta $6.000 por metro de consumo. La receta se expresa en metros, no en kilos.")
fields(doc,[["name","Nombre específico y variante.","Sí","Jersey poliéster blanco"],["category","tela, hilo, avío, accesorio, packaging.","Sí","tela"],["unit","Unidad de compra.","Sí","kilo / metro / unidad"],["unitPrice","Precio de compra actual.","Sí para cotizar","$12.000 por kilo"],["supplierId","Proveedor seleccionado.","No al inicio","Textil SA"],["color","Color/código.","No","Blanco óptico"],["width","Ancho útil.","No","150 cm"],["gramsPerMeter","Gramaje.","No","160 g/m"],["metersPerKilo","Metros útiles por kg.","Condicional","4,5 m/kg"],["yieldPercent","Aprovechamiento operativo.","No","85%"],["notes","Lote, calidad, equivalencias.","No","No mezclar con blanco nieve"]])
heading(doc,"5.1 Separar tres datos",2); bullets(doc,["Rendimiento del material comprado: por ejemplo, 4,5 metros por kilo.","Consumo base de una unidad: por ejemplo, 0,50 metros por camiseta.","Desperdicio de la receta: por ejemplo, 8% por encimado, corte o descarte."])

heading(doc,"6. Catálogo general y clasificación")
body(doc,"Categoría, subcategoría y tipo general son un catálogo editable que aparece por SELECT. Se administra en /admin/configuracion/catalogo. No usar textarea para inventar categorías en cada producto.")
table(doc,["Nivel","Pregunta","Ejemplos","Diferente de"],[["Categoría","¿Gran grupo comercial?","Indumentaria, Mercería, Bolsos y mochilas, Hogar y jardín, Escolar.","Familia."],["Subcategoría","¿Rama del grupo?","Prendas, Conjuntos y kits, Insumos y avíos, Mochilas infantiles.","Tipo de prenda."],["Tipo","¿Qué artículo es?","Remera, Camiseta, Tela, Bolso, Mantel.","garmentType."],["Familia","¿Qué estructura comparte?","Parte superior, Campera, Pantalón, Short fútbol, Bermuda.","Categoría."],["Molde","¿Qué patrón y medidas usa?","Molde superior unisex.","Receta."]],[1.0,2.0,3.0,1.6],7.55)
body(doc,"Valores iniciales: Indumentaria > Prendas (remera, chomba, camiseta, campera, pantalón, short, bermuda, guardapolvo); Indumentaria > Conjuntos y kits; Mercería > Insumos y avíos; Bolsos y mochilas > Bolsos y Mochilas infantiles; Hogar y jardín > Manteles; Escolar > Guardapolvos.")

heading(doc,"7. Familias y moldes")
body(doc,"La familia identifica una estructura y el molde define patrón y medidas requeridas/opcionales. Remera, chomba y camiseta pueden compartir molde, pero siempre son productos distintos y conservan recetas propias. Compartir molde no comparte BOM automáticamente.")
table(doc,["Familia","Casos","Regla"],[["parte_superior","remera, chomba, camiseta","Molde compartible; recetas, SKU y talles separados."],["campera","campera deportiva","Molde propio; validar medidas."],["pantalon","pantalón","Cintura, cadera, tiro y largo."],["short_futbol","short de fútbol","Puede acompañar camiseta; BOM propio."],["bermuda","bermuda","No asumir que es short."],["accesorio","bolsos y otros","Molde/talle opcional; BOM general."]],[1.35,2.0,4.0],8.0)
body(doc,"Campera, pantalón, short de fútbol y bermuda se cargan como productos independientes. Bolsos, mochilas infantiles, manteles y guardapolvos pueden no tener molde o talle; se usa receta general si se cotizan por cantidad directa.")
note(doc,"Invariante molde/familia.","Si se elige un molde, su familia debe ser compatible con la familia del producto. Si no hay certeza, dejar el molde pendiente y preguntar.")

heading(doc,"8. Productos")
body(doc,"Un producto es un artículo vendible: una prenda, un bolso o un conjunto. No es una tela, una familia ni una receta.")
fields(doc,[["sku","Código único estable.","Sí","CAM-FUT-001"],["name","Nombre comercial.","Sí","Camiseta fútbol sublimada"],["description","Diferenciadores.","No","Manga corta, cuello redondo"],["productCategory","Categoría del SELECT.","Sí","indumentaria"],["productSubcategory","Subcategoría.","Sí","prendas"],["productType","Tipo general.","Sí","camiseta"],["productKind","garment o bundle.","Sí","garment"],["garmentFamily","Familia textil.","No","parte_superior"],["garmentType","Tipo de prenda.","No","camiseta"],["moldId","Molde compatible.","No","Molde superior"],["basePrice","Referencia heredada/override.","No","No reemplaza BOM"],["minOrder","Pedido mínimo.","No","6"],["zones","Zonas de aplicación.","No","frente, espalda, manga"]])
table(doc,["productKind","Significado","Condición"],[["garment","Artículo individual.","Puede tener talles, molde y BOM general/por talle."],["bundle","Conjunto/kit.","Componentes garment; nunca bundle anidado."]],[1.1,2.7,3.6],8.1,"5C7897")
note(doc,"Precio.","basePrice queda como referencia. La fuente normal es costo BOM + técnicas + margen + recargo urgente + redondeo. Solo overrideUnitPrice es una excepción manual.")

heading(doc,"9. Talles y medidas")
body(doc,"Los talles pertenecen al producto. Las medidas terminadas son numéricas y se expresan normalmente en centímetros; no se usan para inventar el consumo de corte.")
fields(doc,[["label","Etiqueta comercial.","Sí","6, 8, 10, S, M, L, XL, Único"],["order","Orden visual.","No","10 después de 8"],["measurements","Medidas terminadas.","Según esquema","ancho_pecho=52; largo=70"],["productId","Producto propietario.","Sí","Camiseta fútbol"]])
body(doc,"El molde puede definir como requeridas ancho de pecho y largo, y como opcionales hombro, manga, cuello, cintura, cadera, tiro, boca, fuelle o asas. Si el consumo cambia con el talle, la receta debe ser específica por talle.")
note(doc,"Sin talle.","No crear talles ficticios. Una mochila, un mantel o un accesorio puede usar BOM general con sizeId nulo y cotización por unidades.")

heading(doc,"10. Recetas y BOM")
body(doc,"La receta es la relación entre producto, talle/técnica y materiales. Cada ítem indica método, cantidad o rendimiento y desperdicio.")
fields(doc,[["productId","Producto de la receta.","Sí","Camiseta fútbol"],["sizeId","Talle; nulo si general.","No","M"],["techniqueId","Técnica; nulo si general.","No","Sublimación"],["notes","Corte, versión, supuestos.","No","Encimado de dos"]])
fields(doc,[["materialId","Material existente.","Sí","Jersey blanco"],["consumptionMode","direct o yield.","Sí","yield"],["directQuantity","Consumo por unidad.","Condicional","0,50 m"],["unitsPerConsumptionUnit","Unidades por metro/unidad.","Condicional","2 prendas/m"],["wastePercent","Merma del ítem.","No","8%"]])
table(doc,["Modo","Uso","Fórmula","Ejemplo"],[["direct","Conocemos consumo unitario.","base=directQuantity","0,50 m por camiseta"],["yield","Conocemos rendimiento.","base=1/unitsPerConsumptionUnit","1 m rinde 2 prendas"]],[.9,2.25,2.0,2.25],8.0,"5C7897")
code(doc,"consumo_base = directQuantity\nconsumo_base = 1 / unitsPerConsumptionUnit\nconsumo_total = consumo_base × (1 + wastePercent/100)\nsi unit=kilo: costo_metro = unitPrice / metersPerKilo\ncosto_item = consumo_total × costo_metro\ncosto_prenda = suma(items) + técnicas")
body(doc,"Ejemplo obligatorio: si en un metro entran dos prendas, el consumo unitario es 1/2 metro antes de desperdicio. Con 8% de merma: 0,50 × 1,08 = 0,54 m. Con $12.000/kg y 4,5 m/kg, el metro cuesta $2.666,67 y el ítem cuesta aproximadamente $1.440.")
note(doc,"Bloqueo.","Sin receta usable no se cotiza. Con unit=kilo y sin metersPerKilo tampoco. El error debe indicar completar /admin/recetas o /admin/insumos.")

heading(doc,"11. Técnicas, precios y reglas")
body(doc,"Una técnica separa el costo de trabajo o aplicación: sublimación, DTF, bordado, serigrafía, vinilo u otra.")
fields(doc,[["name","Nombre.","Sí","Sublimación"],["costPerUnit","Costo por unidad.","No","$1.200"],["costPerSquareMeter","Costo por m².","No","$2.500"],["setupCost","Preparación una vez por línea/lote.","No","$4.000"],["description","Incluye/no incluye.","No","Impresión y planchado"],["active","Disponible.","No","Sí"]])
body(doc,"Una técnica por m² exige aplicación en la línea con técnica, zona, ancho y alto. Área = ancho × alto; sin medidas se bloquea y no se inventa.")
code(doc,"costo = materiales + técnica por unidad + setup\nprecio = costo × (1 + margen/100)\nsi urgente: precio × (1 + recargo/100)\nfinal = Math.round(precio/redondeo) × redondeo")
table(doc,["Regla","Función","Ejemplo"],[["marginPercent","Margen sobre costo.","40%"],["urgentSurcharge","Recargo urgente.","15%"],["minAdvancePercent","Seña mínima.","50%"],["rounding","Paso comercial.","100"]],[1.7,3.5,2.0],8.1)

heading(doc,"12. Conjuntos y kits")
body(doc,"Un bundle agrupa productos garment existentes con cantidad y estrategia de talle. El sistema suma el costo de cada componente; no crea una receta única implícita.")
table(doc,["Dato","Significado","Ejemplo"],[["componentProduct","Producto individual.","Camiseta fútbol"],["quantity","Cantidad incluida.","1"],["same_label","Mismo rótulo si existe.","M busca M en camiseta/short"],["fixed","Talle fijo del componente.","Talle 8"],["componentSizeId","Talle fijo si fixed.","8"]],[1.65,3.2,2.35],8.0,"5C7897")
bullets(doc,["Camiseta + short de fútbol: dos componentes, cantidad 1, same_label; BOM propios.","Campera + pantalón: same_label si las etiquetas equivalen; si no, fixed o equivalencia explícita.","Bermuda + chomba: same_label solo si se corresponden; BOM de bermuda y chomba separados."])
note(doc,"Prohibición.","No permitir kits anidados. Un bundle no puede ser componente de otro bundle.")

heading(doc,"13. Pedidos, cotización pública, pagos y snapshots")
table(doc,["Elemento","Qué captura","Relación"],[["Organización/contacto","Cliente y referente.","Pedido."],["Línea","Producto, cantidad, talle y técnica.","Catálogo."],["Order item","Talle, nombre, número, stage.","Planilla individual."],["Arte/aplicación","Archivo, zona, técnica, medidas.","Línea/adjunto."],["Pago","Importe, fecha, medio, referencia.","Caja/bloqueo."],["Snapshot","Regla, consumo, precios y totales.","Historial inmutable."]],[1.55,3.2,2.45],7.9)
code(doc,"borrador → presupuesto_enviado → aprobado → seniado → en_produccion\n          → corte → confeccion → estampado → control → entregado\ncualquier estado → bloqueado_pago si pagos < seña mínima\ncualquier estado → cancelado")
body(doc,"La cotización pública crea un borrador con publicToken; el cliente no ve costos internos. El panel valida cantidades, talles, recetas, técnicas y aplicaciones antes de confirmar. quoteOrder calcula en servidor.")
body(doc,"Confirmar o re-cotizar guarda snapshot con regla, líneas, desglose por talle, materiales, técnicas, costos y totales. Una nueva cotización agrega historial y no pisa la anterior. Los pagos cancelados quedan visibles, pero no cuentan.")

heading(doc,"14. Política flexible e invariantes")
body(doc,"Faltantes opcionales no bloquean: dirección, CUIT, color, gramaje, medidas opcionales, molde de un no textil, descripción, zonas o técnica aún no decidida. Se marca pendiente.")
table(doc,["Invariante","Por qué","Qué hacer"],[["SKU, nombre, categoría y productKind","Identidad base.","Preguntar antes de guardar."],["bundle con componentes","No se cotiza vacío.","Cargar garment components."],["Sin bundle anidado","Evita recursión.","Desarmar/revisar."],["Molde y familia compatibles","Medidas coherentes.","Dejar molde pendiente."],["Números válidos","Evita cálculos inválidos.","Pedir valor y unidad."],["metersPerKilo si kilo","Costo confiable.","Completar material."],["Receta cotizable","Evita costo cero.","Completar BOM."],["Talle pertenece al producto","Evita relaciones cruzadas.","Seleccionar talle correcto."]],[1.7,3.6,1.9],7.55,"5C7897")

heading(doc,"15. Mapa de páginas y APIs")
table(doc,["Ruta","Qué captura","Momento"],[["/admin/proveedores","Proveedor y contacto.","Antes de materiales."],["/admin/insumos","Material, unidad, precio, proveedor y rendimiento.","Cada compra."],["/admin/tecnicas","Costos por unidad/m² y setup.","Antes de BOM."],["/admin/configuracion/catalogo","Categoría, subcategoría y tipo por SELECT.","Cuando falta una opción."],["/admin/moldes","Familia, molde y medidas.","Textiles."],["/admin/productos/nuevo","Producto garment/bundle.","Alta."],["/admin/productos/[id]/talles","Talles/medidas.","Venta por talle."],["/admin/productos/[id]/recetas","Recetas e ítems BOM.","Consumo y costo."],["/admin/recetas","Vista global/cobertura.","Control."],["/admin/pedidos/[id]/cotizar","Preview/confirmar/re-cotizar.","Con BOM listo."],["/admin/pedidos/[id]/pagos","Pagos y seña.","Desbloqueo."],["/admin/caja","Movimientos.","Finanzas."],["/admin/pedidos/[id]/ficha-tecnica","Consumos, técnicas y arte.","Producción."]],[2.35,3.25,1.6],7.2)
body(doc,"APIs principales: GET/POST /api/suppliers; PATCH /api/suppliers/[id]; POST/PATCH /api/materials; POST/PATCH /api/techniques; GET/POST /api/product-taxonomy; POST/PATCH /api/products; /api/products/[id]/sizes; /api/garment-molds; /api/recipes y /api/recipes/[id]/items; /api/orders, /api/order-lines, /api/order-items y /api/payments.")

heading(doc,"16. Plantillas para audio")
for title,text in [("16.1 Proveedor","PROVEEDOR\nNombre:\nTeléfono / WhatsApp:\nEmail:\nContacto:\nDirección/localidad:\nCUIT:\nQué vende:\nMínimos, plazos, notas:\nActivo: sí / no / después"),("16.2 Insumo","INSUMO\nNombre y variante:\nCategoría: tela / hilo / avío / accesorio / packaging\nUnidad: kilo / metro / unidad / rollo / pack\nPrecio y fecha:\nProveedor:\nColor / ancho / gramaje:\nSi es kilo, metros por kilo:\nRendimiento o desperdicio:\nNotas:"),("16.3 Producto","PRODUCTO\nSKU:\nNombre:\nDescripción:\nCategoría / subcategoría / tipo:\ngarment o bundle:\nFamilia / tipo de prenda:\nMolde:\nTalles:\nPedido mínimo:\nZonas:\nPendientes:"),("16.4 Receta","RECETA\nProducto:\nTalle: general / 6 / 8 / S / M / L / otro\nTécnica:\nMaterial:\nModo direct o yield:\nCantidad o rendimiento:\nDesperdicio %:\nObservaciones:"),("16.5 Conjunto","CONJUNTO\nSKU y nombre:\nComponente 1 + cantidad:\nComponente 2 + cantidad:\nMapping: same_label o fixed:\nTalle fijo:\nConfirmar: ningún componente es bundle")]:
    heading(doc,title,2); code(doc,text)
heading(doc,"16.6 Checklist antes de cotizar",2); bullets(doc,["Producto con SKU, nombre, categoría y productKind.","Bundle con componentes garment y sin kit anidado.","Receta general o por talle.","Materiales existentes con unidad.","metersPerKilo en cada material kilo.","Técnica y medidas para cada aplicación por m².","Cantidades y talles desglosados.","Snapshot al confirmar."],True)

heading(doc,"17. Gobierno de datos y trazabilidad")
table(doc,["Situación","Acción correcta","No hacer"],[["Sube precio de tela","Editar unitPrice y registrar fuente/fecha.","Modificar snapshot."],["Cambia proveedor","Actualizar supplierId o crear variante.","Texto ambiguo."],["Material discontinuado","active=false; conservar historial.","Borrar referenciado."],["Corrige talle","Preservar IDs; retirar solo sin referencias.","Borrar/recrear sin control."],["Cambia receta","Editar para futuro y registrar motivo.","Recalcular confirmado."],["Cambia margen","Actualizar regla vigente.","Reescribir compromiso."],["Reemplaza arte","Nueva versión/aprobación.","Sobrescribir evidencia."]],[1.6,3.9,1.8],7.5,"5C7897")
bullets(doc,["SKU: prefijo + variante + número estable, por ejemplo CAM-FUT-001.","Material: tipo + composición/uso + color + variante.","Notas: fecha, fuente del precio, rendimiento y supuestos.","Preferir active=false para no romper relaciones.","Registrar quién, qué, por qué y desde cuándo en cambios sensibles."])

heading(doc,"18. Estado técnico actual")
table(doc,["Tema","Estado documentado"],[["Neon","Migraciones 0003 a 0008 aplicadas: proveedores, integridad, BOM, familias/moldes/conjuntos y catálogo editable."],["Seed","No se ejecutó seed; no asumir datos demo."],["Verificaciones","Typecheck, lint y build verificados; puede haber warnings preexistentes de lint."],["Validación","Zod; proveedor normalizado; kilo exige metersPerKilo; receta valida talle/producto."],["Pricing","quoteOrder server-side con BOM, conversión kilo/metro y redondeo."],["Historial","snapshot y snapshotHistory conservan contexto; cambios futuros no los pisan."],["Taxonomía","Se administra en /admin/configuracion/catalogo y se usa por SELECT."],["Seguridad","No se incluyen secretos ni URL de DB."]],[1.55,5.75],7.65)
note(doc,"Límite conocido.","Una técnica por m² requiere aplicación con técnica, ancho y alto. Sin esos datos, la cotización devuelve un error accionable. materials.supplier existe por compatibilidad; supplierId es canónico.")

doc.add_page_break(); heading(doc,"Apéndice. Ejemplos compactos")
table(doc,["Registro","Datos","Relación"],[["Proveedor","Textil SA; +54 370 4XX XXXX; ventas@textil.example; contacto María López","supplierId en materiales."],["Tela","Jersey poliéster blanco; tela; kilo; $12.000; 4,5 m/kg; ancho 150","Textil SA; BOM."],["Hilo","Hilo overlock blanco; hilo; unidad; $3.500","Hilos del Norte; direct."],["Avío","Ribb blanco cuello; avío; metro; $1.800","Textil SA; BOM variable."],["Packaging","Bolsa mediana; packaging; unidad; $120","BOM general: 1 por producto."]],[1.2,4.35,1.75],7.5,"5C7897")
table(doc,["Entidad","Ejemplo completo"],[["Producto","CAM-FUT-001; Camiseta fútbol sublimada; indumentaria/prendas/camiseta; garment; parte_superior; camiseta; molde superior; mínimo 6; zonas frente/espalda/manga."],["Talles","S, M, L, XL con medidas numéricas; IDs preservados al editar."],["BOM tela","Jersey; yield; 2 prendas/m; 8%; 0,50 m base y 0,54 m total."],["BOM hilo","Hilo; direct; 12 m o unidad documentada; 5%."],["Técnica","Sublimación; $1.200 unidad o aplicación con ancho/alto si m²."],["Precio","Costo + margen 40%; urgente 15%; redondeo 100."]],[1.25,6.05],7.55)
table(doc,["Bundle","Componentes"],[["SET-FUT-001","Conjunto fútbol completo; bundle; camiseta cantidad 1 same_label + short cantidad 1 same_label."],["Regla","Cada componente tiene talles/BOM propios; nunca incluir otro bundle."]],[1.45,5.85],7.7,"5C7897")
code(doc,"Textil SA\n  └─ Jersey blanco (kilo, $12.000, 4,5 m/kg)\n       └─ BOM CAM-FUT-001 / M / yield 2 por metro / merma 8%\n            └─ CAM-FUT-001 (garment, parte_superior, molde superior)\n                 └─ SET-FUT-001 (bundle)\n                      ├─ CAM-FUT-001 (same_label, 1)\n                      └─ SHO-FUT-001 (same_label, 1)\n                           └─ Pedido → cotización → pago → snapshot → producción")
heading(doc,"Preguntas de cierre",2); bullets(doc,["¿Es una variante real o mezclamos materiales?","¿Unidad y precio son de compra?","¿Proveedor seleccionado o pendiente?","¿Individual o conjunto?","¿Categoría del SELECT?","¿Molde y familia compatibles?","¿Consumo en unidad correcta?","¿Merma separada del rendimiento?","¿Puede cotizarse sin inventar?","¿Qué queda pendiente?"],True)
note(doc,"Próximo uso.","Transcribir cada audio por bloques de proveedor, material, producto, talle, técnica, receta o conjunto. Validar relaciones y cargar siguiendo esta guía.")

doc.core_properties.title="Guía maestra de carga de datos AH Sports"; doc.core_properties.subject="Manual operativo de datos y cotización en Neon"; doc.core_properties.author="AH Sports"; doc.core_properties.keywords="AH Sports, proveedores, insumos, productos, recetas, BOM, Neon"
OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT); print(str(OUT))
